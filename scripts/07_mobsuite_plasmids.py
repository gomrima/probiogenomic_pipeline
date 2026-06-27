#!/usr/bin/env python3
"""
Memory-safe MOB-suite plasmid module.

Runs MOB-suite mob_recon on each normalized genome FASTA and consolidates the
per-genome contig reports into the standardized pipeline outputs.

Why this rewrite
----------------
The previous implementation accumulated every per-genome contig report in a list
of pandas DataFrames, called pandas.concat() on all of them at once, and then
wrote the full detail table to XLSX with the in-memory openpyxl engine. On an
8 GB machine this could saturate RAM during the final consolidation.

This version keeps identical outputs but:
  - streams detail rows straight to mobsuite_contig_detail.tsv (no concat);
  - computes the per-genome summary and the binary feature set while streaming;
  - does not capture the mob_recon subprocess output in Python (it is tee'd to
    the module log by the .sh wrapper);
  - writes the XLSX workbook with openpyxl in constant-memory (write_only) mode.

Output filenames are preserved:
  - mobsuite_contig_detail.tsv
  - mobsuite_feature_binary.tsv   (read by module 17)
  - mobsuite_summary_by_genome.tsv
  - mobsuite_failures.tsv
  - mobsuite_results.xlsx         (human-facing convenience workbook)
"""

from __future__ import annotations

import argparse
import csv
import gc
import os
import shutil
import subprocess
import sys
import re
from pathlib import Path
from typing import Dict, Iterable, List, Tuple

csv.field_size_limit(min(sys.maxsize, 2_147_483_647))


# Columns written by MOB-suite in contig_report.txt (used for header validation
# and for the detail table, after the three provenance columns are prepended).
MOBSUITE_REQUIRED_COLUMNS = [
    "sample_id", "molecule_type", "primary_cluster_id", "secondary_cluster_id",
    "contig_id", "size", "gc", "md5", "circularity_status", "rep_type(s)",
    "rep_type_accession(s)", "relaxase_type(s)", "relaxase_type_accession(s)",
    "mpf_type", "mpf_type_accession(s)", "orit_type(s)", "orit_accession(s)",
    "predicted_mobility", "mash_nearest_neighbor", "mash_neighbor_distance",
    "mash_neighbor_identification", "repetitive_dna_id", "repetitive_dna_type",
    "filtering_reason",
]

DETAIL_COLUMNS = ["genome_id", "input_fna", "result_file"] + MOBSUITE_REQUIRED_COLUMNS

SUMMARY_COLUMNS = [
    "genome_id",
    "input_fna",
    "n_contigs_reported",
    "n_plasmid_contigs",
    "n_chromosome_contigs",
    "n_unique_primary_clusters",
    "n_unique_secondary_clusters",
    "n_rep_types",
    "n_relaxase_types",
    "n_mpf_types",
    "n_orit_types",
    "n_mobility_classes",
]

FAIL_COLUMNS = [
    "genome_id",
    "input_fna",
    "status",
    "message",
]

MANIFEST_REQUIRED_COLUMNS = ["genome_id", "normalized_fna", "status"]

EMPTY_TOKENS = {"", "-", "--", "nan", "NaN", "None"}

EXCEL_MAX_ROWS = 1_048_576


def safe_mkdir(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)


def sanitize_name(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", str(name))


def clean_cell(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def run_command(cmd: List[str], log_prefix: str = "") -> None:
    """Run a command without buffering its stdout/stderr in Python memory."""
    print(f"{log_prefix}Running: {' '.join(cmd)}", flush=True)
    result = subprocess.run(cmd)
    if result.returncode != 0:
        raise RuntimeError(
            f"Command failed with exit code {result.returncode}: {' '.join(cmd)} "
            f"(see module log for the full mob_recon output)"
        )


def read_manifest_lowmem(manifest_path: Path) -> List[Dict[str, str]]:
    rows: List[Dict[str, str]] = []
    with manifest_path.open("r", newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        fieldnames = reader.fieldnames or []
        missing = [c for c in MANIFEST_REQUIRED_COLUMNS if c not in fieldnames]
        if missing:
            raise ValueError(f"Missing required manifest columns: {missing}")
        for row in reader:
            if clean_cell(row.get("status")) != "OK":
                continue
            fna = clean_cell(row.get("normalized_fna"))
            if not fna or not Path(fna).is_file():
                continue
            rows.append({
                "genome_id": clean_cell(row.get("genome_id")),
                "normalized_fna": fna,
            })
    return rows


def clean_token_set(field_value: object) -> List[str]:
    s = clean_cell(field_value)
    if s in EMPTY_TOKENS:
        return []
    cleaned = []
    for token in re.split(r"[,\|;]+", s):
        token = token.strip()
        if token in EMPTY_TOKENS:
            continue
        cleaned.append(token)
    return sorted(set(cleaned))


def write_tsv(path: Path, columns: List[str], rows: Iterable[Dict[str, object]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, delimiter="\t", extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def write_feature_binary_matrix(
    path: Path,
    genome_order: List[str],
    genome_features: Dict[str, set],
) -> None:
    all_features = sorted({feature for features in genome_features.values() for feature in features})
    columns = ["genome_id"] + all_features
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, delimiter="\t")
        writer.writeheader()
        for genome_id in genome_order:
            present = genome_features.get(genome_id, set())
            row = {"genome_id": genome_id}
            for feature in all_features:
                row[feature] = 1 if feature in present else 0
            writer.writerow(row)


def _stream_tsv_into_sheet(ws, tsv_path: Path, max_rows: int = EXCEL_MAX_ROWS) -> bool:
    truncated = False
    with tsv_path.open("r", newline="", encoding="utf-8", errors="replace") as handle:
        reader = csv.reader(handle, delimiter="\t")
        for i, record in enumerate(reader):
            if i >= max_rows:
                truncated = True
                break
            ws.append(record)
    return truncated


def write_workbook_lowmem(xlsx_path: Path, sheet_specs: List[Tuple[str, Path]]) -> None:
    """Write the workbook in constant memory (openpyxl write_only), streaming each TSV."""
    try:
        from openpyxl import Workbook
    except Exception as exc:  # pragma: no cover - environment guard
        note = xlsx_path.with_suffix(xlsx_path.suffix + ".NOT_CREATED.txt")
        note.write_text(
            f"XLSX workbook was not created because openpyxl could not be imported: {exc}\n"
            f"All TSV outputs remain valid and complete.\n",
            encoding="utf-8",
        )
        print(f"[WARN] XLSX not created. See: {note}", flush=True)
        return

    workbook = Workbook(write_only=True)
    truncated_sheets: List[str] = []
    for title, tsv_path in sheet_specs:
        ws = workbook.create_sheet(title=title[:31])
        if tsv_path.is_file():
            if _stream_tsv_into_sheet(ws, tsv_path):
                truncated_sheets.append(title)
        else:
            ws.append(["(missing TSV)", str(tsv_path)])

    if truncated_sheets:
        ws = workbook.create_sheet(title="xlsx_notes")
        ws.append(["sheet", "message"])
        for sheet in truncated_sheets:
            ws.append([sheet, "Truncated at the Excel row limit; the complete data are in the corresponding TSV."])

    workbook.save(str(xlsx_path))


def write_empty_outputs(output_tsv_dir: Path, output_xlsx_dir: Path) -> None:
    detail_out = output_tsv_dir / "mobsuite_contig_detail.tsv"
    binary_out = output_tsv_dir / "mobsuite_feature_binary.tsv"
    summary_out = output_tsv_dir / "mobsuite_summary_by_genome.tsv"
    fail_out = output_tsv_dir / "mobsuite_failures.tsv"
    xlsx_out = output_xlsx_dir / "mobsuite_results.xlsx"

    write_tsv(detail_out, DETAIL_COLUMNS, [])
    write_tsv(summary_out, SUMMARY_COLUMNS, [])
    write_tsv(fail_out, FAIL_COLUMNS, [])
    write_tsv(binary_out, ["genome_id"], [])
    write_workbook_lowmem(
        xlsx_out,
        [
            ("mobsuite_contig_detail", detail_out),
            ("mobsuite_feature_binary", binary_out),
            ("mobsuite_summary", summary_out),
            ("mobsuite_failures", fail_out),
        ],
    )


def process_contig_report_streaming(
    report_path: Path,
    genome_id: str,
    input_fna: Path,
    detail_writer: "csv.DictWriter",
) -> Tuple[Dict[str, object], set]:
    """Stream a contig_report.txt to the detail TSV and return (summary, features).

    All per-genome aggregates are computed on the fly so no genome-wide DataFrame
    is ever held in memory.
    """
    n_contigs = 0
    n_plasmid = 0
    n_chromosome = 0
    primary_clusters: set = set()
    secondary_clusters: set = set()
    rep_types: set = set()
    relaxase_types: set = set()
    mpf_types: set = set()
    orit_types: set = set()
    mobility_classes: set = set()
    features: set = set()

    with report_path.open("r", newline="", encoding="utf-8", errors="replace") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        fieldnames = reader.fieldnames or []
        missing = [c for c in MOBSUITE_REQUIRED_COLUMNS if c not in fieldnames]
        if missing:
            raise ValueError(f"Missing expected MOB-suite contig_report columns: {missing}")

        for row in reader:
            out = {col: "" for col in DETAIL_COLUMNS}
            out["genome_id"] = genome_id
            out["input_fna"] = str(input_fna)
            out["result_file"] = str(report_path)
            for col in MOBSUITE_REQUIRED_COLUMNS:
                out[col] = clean_cell(row.get(col))
            detail_writer.writerow(out)

            n_contigs += 1
            molecule_type = out["molecule_type"].lower()

            if molecule_type == "plasmid":
                n_plasmid += 1
                features.add("plasmid_present")

                primary = out["primary_cluster_id"]
                secondary = out["secondary_cluster_id"]
                if primary not in EMPTY_TOKENS:
                    primary_clusters.add(primary)
                if secondary not in EMPTY_TOKENS:
                    secondary_clusters.add(secondary)

                for token in clean_token_set(out["rep_type(s)"]):
                    rep_types.add(token)
                    features.add(f"replicon:{token}")
                for token in clean_token_set(out["relaxase_type(s)"]):
                    relaxase_types.add(token)
                    features.add(f"relaxase:{token}")
                for token in clean_token_set(out["mpf_type"]):
                    mpf_types.add(token)
                    features.add(f"mpf:{token}")
                for token in clean_token_set(out["orit_type(s)"]):
                    orit_types.add(token)
                    features.add(f"orit:{token}")
                for token in clean_token_set(out["predicted_mobility"]):
                    mobility_classes.add(token)
                    features.add(f"mobility:{token}")

            elif molecule_type == "chromosome":
                n_chromosome += 1

    summary = {
        "genome_id": genome_id,
        "input_fna": str(input_fna),
        "n_contigs_reported": n_contigs,
        "n_plasmid_contigs": n_plasmid,
        "n_chromosome_contigs": n_chromosome,
        "n_unique_primary_clusters": len(primary_clusters),
        "n_unique_secondary_clusters": len(secondary_clusters),
        "n_rep_types": len(rep_types),
        "n_relaxase_types": len(relaxase_types),
        "n_mpf_types": len(mpf_types),
        "n_orit_types": len(orit_types),
        "n_mobility_classes": len(mobility_classes),
    }
    return summary, features


def zero_summary(genome_id: str, input_fna: Path) -> Dict[str, object]:
    return {
        "genome_id": genome_id,
        "input_fna": str(input_fna),
        "n_contigs_reported": 0,
        "n_plasmid_contigs": 0,
        "n_chromosome_contigs": 0,
        "n_unique_primary_clusters": 0,
        "n_unique_secondary_clusters": 0,
        "n_rep_types": 0,
        "n_relaxase_types": 0,
        "n_mpf_types": 0,
        "n_orit_types": 0,
        "n_mobility_classes": 0,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Run MOB-suite mob_recon on normalized genome FASTA inputs and consolidate plasmid outputs.")
    parser.add_argument("--manifest", required=True, help="Path to input_manifest.tsv")
    parser.add_argument("--raw-root", required=True, help="Root directory for per-genome MOB-suite raw outputs")
    parser.add_argument("--db-dir", required=True, help="Path to MOB-suite database directory")
    parser.add_argument("--output-tsv-dir", required=True, help="Directory for TSV outputs")
    parser.add_argument("--output-xlsx-dir", required=True, help="Directory for XLSX outputs")
    parser.add_argument("--threads", type=int, default=2, help="Max threads for mob_recon")
    args = parser.parse_args()

    manifest_path = Path(args.manifest).resolve()
    raw_root = Path(args.raw_root).resolve()
    db_dir = Path(args.db_dir).resolve()
    output_tsv_dir = Path(args.output_tsv_dir).resolve()
    output_xlsx_dir = Path(args.output_xlsx_dir).resolve()
    threads = max(1, min(args.threads, 2))

    safe_mkdir(raw_root)
    safe_mkdir(output_tsv_dir)
    safe_mkdir(output_xlsx_dir)

    if not manifest_path.is_file():
        raise FileNotFoundError(f"Manifest not found: {manifest_path}")
    if not db_dir.is_dir():
        raise FileNotFoundError(f"MOB-suite database directory not found: {db_dir}")

    manifest_rows = read_manifest_lowmem(manifest_path)

    if not manifest_rows:
        print("No valid FNA entries found in manifest. Writing empty outputs.", flush=True)
        write_empty_outputs(output_tsv_dir, output_xlsx_dir)
        return

    detail_out = output_tsv_dir / "mobsuite_contig_detail.tsv"
    binary_out = output_tsv_dir / "mobsuite_feature_binary.tsv"
    summary_out = output_tsv_dir / "mobsuite_summary_by_genome.tsv"
    fail_out = output_tsv_dir / "mobsuite_failures.tsv"
    xlsx_out = output_xlsx_dir / "mobsuite_results.xlsx"

    # Consolidated outputs are written to temp files and published atomically
    # (os.replace) only after the full pass, so an interruption never leaves a
    # half-written TSV for the downstream module 17.
    detail_tmp = detail_out.with_name(detail_out.name + ".tmp")
    binary_tmp = binary_out.with_name(binary_out.name + ".tmp")
    summary_tmp = summary_out.with_name(summary_out.name + ".tmp")
    fail_tmp = fail_out.with_name(fail_out.name + ".tmp")
    xlsx_tmp = xlsx_out.with_name(xlsx_out.name + ".tmp")

    summary_records: List[Dict[str, object]] = []
    failure_records: List[Dict[str, object]] = []
    genome_features: Dict[str, set] = {}
    genome_order: List[str] = []

    detail_handle = detail_tmp.open("w", newline="", encoding="utf-8")
    detail_writer = csv.DictWriter(detail_handle, fieldnames=DETAIL_COLUMNS, delimiter="\t", extrasaction="ignore")
    detail_writer.writeheader()

    try:
        for i, row in enumerate(manifest_rows, start=1):
            genome_id = row["genome_id"]
            input_fna = Path(row["normalized_fna"]).resolve()
            run_dir = raw_root / sanitize_name(genome_id)

            genome_order.append(genome_id)
            genome_features.setdefault(genome_id, set())

            print(f"\n=== [{i}/{len(manifest_rows)}] Processing genome: {genome_id} ===", flush=True)
            print(f"FNA: {input_fna}", flush=True)
            print(f"MOB-suite run directory: {run_dir}", flush=True)

            try:
                # mob_recon refuses to write into an existing directory, so a
                # stale per-genome directory from an interrupted run would make
                # every rerun fail. Remove it first so reruns start clean.
                if run_dir.exists():
                    shutil.rmtree(run_dir)

                cmd = [
                    "mob_recon",
                    "-i", str(input_fna),
                    "-o", str(run_dir),
                    "-d", str(db_dir),
                    "-n", str(threads),
                ]
                run_command(cmd, log_prefix=f"[{genome_id}] ")

                report_path = run_dir / "contig_report.txt"
                if not report_path.is_file():
                    raise FileNotFoundError(f"Expected contig_report.txt not found: {report_path}")

                summary_record, features = process_contig_report_streaming(
                    report_path, genome_id, input_fna, detail_writer
                )
                summary_records.append(summary_record)
                genome_features[genome_id] = features

            except Exception as exc:
                failure_records.append({
                    "genome_id": genome_id,
                    "input_fna": str(input_fna),
                    "status": "FAILED",
                    "message": str(exc),
                })
                summary_records.append(zero_summary(genome_id, input_fna))
                print(f"[ERROR] Genome failed: {genome_id}", flush=True)
                print(str(exc), flush=True)

            gc.collect()

    finally:
        detail_handle.close()

    summary_records_sorted = sorted(summary_records, key=lambda x: str(x.get("genome_id", "")))
    failure_records_sorted = sorted(failure_records, key=lambda x: str(x.get("genome_id", "")))

    write_tsv(summary_tmp, SUMMARY_COLUMNS, summary_records_sorted)
    write_tsv(fail_tmp, FAIL_COLUMNS, failure_records_sorted)
    write_feature_binary_matrix(binary_tmp, sorted(genome_order), genome_features)

    # Atomic publish: move every temp TSV into its final name.
    os.replace(detail_tmp, detail_out)
    os.replace(binary_tmp, binary_out)
    os.replace(summary_tmp, summary_out)
    os.replace(fail_tmp, fail_out)

    try:
        write_workbook_lowmem(
            xlsx_tmp,
            [
                ("mobsuite_contig_detail", detail_out),
                ("mobsuite_feature_binary", binary_out),
                ("mobsuite_summary", summary_out),
                ("mobsuite_failures", fail_out),
            ],
        )
        os.replace(xlsx_tmp, xlsx_out)
    except Exception as exc:
        print(f"[WARN] XLSX generation failed (non-fatal): {exc}", flush=True)

    print("\nMOB-suite plasmid module completed (memory-safe).", flush=True)
    print(f"Detail TSV   : {detail_out}", flush=True)
    print(f"Binary TSV   : {binary_out}", flush=True)
    print(f"Summary TSV  : {summary_out}", flush=True)
    print(f"Failures TSV : {fail_out}", flush=True)
    print(f"Workbook XLSX: {xlsx_out}", flush=True)


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"[FATAL] {exc}", file=sys.stderr)
        sys.exit(1)
