#!/usr/bin/env python3
"""
Module 12 - CAZyme screening via hmmscan against the dbCAN HMM database.

Modes
-----
Normal  : --manifest / --dbcan-db / --raw-dir  -> runs hmmscan for each genome
          (skips genomes whose .domtblout already exists in --raw-dir)
Resume  : --resume                              -> parses existing .domtblout files
          without re-running hmmscan
"""

import argparse
import re
import subprocess
import sys
from pathlib import Path

import pandas as pd


# ------------------------------------------------------------
# Utilities
# ------------------------------------------------------------

def sanitize_feature_name(value: str) -> str:
    if value is None:
        return ""
    value = str(value)
    value = value.replace("\t", " ").replace("\r", " ").replace("\n", " ")
    return re.sub(r"\s+", " ", value).strip()


# Characters that are illegal in XML 1.0 (and therefore crash openpyxl)
_ILLEGAL_XML_RE = re.compile(
    r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x84\x86-\x9f\ud800-\udfff\ufffe\uffff]"
)


def _clean_str(val):
    """Remove XML-illegal characters from a string value."""
    if isinstance(val, str):
        return _ILLEGAL_XML_RE.sub("", val)
    return val


def _clean_float(val):
    """Replace inf / -inf with None so openpyxl can serialise the cell."""
    import math
    if isinstance(val, float) and (math.isinf(val) or math.isnan(val)):
        return None
    return val


def clean_df_for_excel(df: pd.DataFrame) -> pd.DataFrame:
    """Return a copy of df safe to write with openpyxl (no illegal XML chars, no inf)."""
    df = df.copy()
    for col in df.columns:
        if df[col].dtype == object:
            df[col] = df[col].map(_clean_str)
        elif df[col].dtype in (float,):
            df[col] = df[col].map(_clean_float)
    return df


def safe_text(value) -> str:
    if pd.isna(value):
        return ""
    return str(value).strip()


def ensure_dir(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)


def run_command(cmd):
    return subprocess.run(
        cmd,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
    )


# ------------------------------------------------------------
# Parsing
# ------------------------------------------------------------

def parse_domtblout(domtbl_path: Path, genome_id: str, input_faa: str) -> pd.DataFrame:
    """Parse an hmmscan --domtblout file and return a detail DataFrame."""
    rows = []

    if not domtbl_path.exists():
        return pd.DataFrame()

    with domtbl_path.open("r", encoding="utf-8", errors="replace") as handle:
        for line in handle:
            line = line.strip()
            if not line or line.startswith("#"):
                continue

            parts = line.split()
            if len(parts) < 23:
                continue

            try:
                target_name      = sanitize_feature_name(parts[0])   # CAZy family (HMM)
                target_accession = parts[1]
                target_len       = int(parts[2])                       # HMM model length

                query_name       = sanitize_feature_name(parts[3])   # protein id
                query_accession  = parts[4]
                query_len        = int(parts[5])

                full_evalue  = float(parts[6])
                full_score   = float(parts[7])
                full_bias    = float(parts[8])

                domain_number = int(parts[9])
                domain_of     = int(parts[10])

                c_evalue     = float(parts[11])
                i_evalue     = float(parts[12])
                domain_score = float(parts[13])
                domain_bias  = float(parts[14])

                hmm_from = int(parts[15])
                hmm_to   = int(parts[16])
                ali_from = int(parts[17])
                ali_to   = int(parts[18])
                env_from = int(parts[19])
                env_to   = int(parts[20])
                acc      = float(parts[21])

                # dbCAN standard filter: HMM model coverage a 35 % and i-Evalue a 1e-15
                hmm_cov = (hmm_to - hmm_from + 1) / target_len if target_len > 0 else 0.0
                passed  = int(i_evalue <= 1e-15 and hmm_cov >= 0.35)

            except (ValueError, IndexError):
                continue

            rows.append(
                {
                    "genome_id":         genome_id,
                    "input_faa":         input_faa,
                    "protein_id":        query_name,
                    "protein_accession": query_accession,
                    "protein_length":    query_len,
                    "cazy_family":       target_name,
                    "family_accession":  target_accession,
                    "family_length":     target_len,
                    "full_evalue":       full_evalue,
                    "full_score":        full_score,
                    "full_bias":         full_bias,
                    "domain_number":     domain_number,
                    "domain_of":         domain_of,
                    "c_evalue":          c_evalue,
                    "i_evalue":          i_evalue,
                    "domain_score":      domain_score,
                    "domain_bias":       domain_bias,
                    "hmm_from":          hmm_from,
                    "hmm_to":            hmm_to,
                    "ali_from":          ali_from,
                    "ali_to":            ali_to,
                    "env_from":          env_from,
                    "env_to":            env_to,
                    "acc":               acc,
                    "hmm_coverage":      round(hmm_cov, 6),
                    "passed_filter":     passed,
                }
            )

    return pd.DataFrame(rows) if rows else pd.DataFrame()


# ------------------------------------------------------------
# Matrix / summary builders
# ------------------------------------------------------------

def build_binary_matrix(detail_df: pd.DataFrame, manifest_df: pd.DataFrame) -> pd.DataFrame:
    genomes = manifest_df["genome_id"].astype(str).tolist()
    base    = pd.DataFrame({"genome_id": genomes})

    if detail_df.empty:
        return base

    pass_df = detail_df.loc[detail_df["passed_filter"] == 1].copy()
    if pass_df.empty:
        return base

    pass_df["cazy_family"] = pass_df["cazy_family"].map(sanitize_feature_name)

    mat = (
        pass_df.assign(value=1)
        .pivot_table(
            index="genome_id",
            columns="cazy_family",
            values="value",
            aggfunc="max",
            fill_value=0,
        )
        .reset_index()
    )
    mat.columns.name = None

    feature_cols = sorted(c for c in mat.columns if c != "genome_id")
    mat = mat[["genome_id"] + feature_cols]
    mat = base.merge(mat, on="genome_id", how="left").fillna(0)

    for col in mat.columns:
        if col != "genome_id":
            mat[col] = pd.to_numeric(mat[col], errors="coerce").fillna(0).astype(int)

    return mat


def build_summary(detail_df: pd.DataFrame, manifest_df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for _, row in manifest_df.iterrows():
        genome_id = str(row["genome_id"])
        input_faa = safe_text(row.get("normalized_faa", ""))

        if not detail_df.empty and "genome_id" in detail_df.columns:
            sub    = detail_df.loc[detail_df["genome_id"] == genome_id]
            passed = sub.loc[sub["passed_filter"] == 1]
        else:
            sub    = pd.DataFrame()
            passed = pd.DataFrame()

        rows.append(
            {
                "genome_id":              genome_id,
                "input_faa":              input_faa,
                "n_total_raw_hits":       int(len(sub)),
                "n_pass_hits":            int(len(passed)),
                "n_unique_pass_proteins": int(passed["protein_id"].nunique()) if not passed.empty else 0,
                "n_unique_pass_families": int(passed["cazy_family"].nunique())  if not passed.empty else 0,
            }
        )
    return pd.DataFrame(rows)


# ------------------------------------------------------------
# Excel writer robust against empty frames
# ------------------------------------------------------------

def _excel_cell(val):
    """Make one cell value safe for openpyxl, matching the old to_excel behavior.

    Reuses the existing per-value cleaners: inf / NaN floats become None (blank
    cell) and XML-illegal characters are stripped from strings. Ints and valid
    floats keep their native type.
    """
    return _clean_str(_clean_float(val))


def _append_df_to_sheet(ws, df: pd.DataFrame, fallback_msg: str) -> None:
    """Stream a DataFrame into a write_only worksheet (constant memory).

    Mirrors the previous per-sheet behavior: a non-empty frame writes its header
    plus its rows (cell types preserved), an empty frame writes a one-row notice.
    """
    if df is not None and not df.empty:
        ws.append([str(c) for c in df.columns])
        for record in df.itertuples(index=False, name=None):
            ws.append([_excel_cell(v) for v in record])
    else:
        ws.append(["status"])
        ws.append([fallback_msg])


def write_xlsx(
    xlsx_path: Path,
    detail_df: pd.DataFrame,
    binary_df: pd.DataFrame,
    summary_df: pd.DataFrame,
    failures_df: pd.DataFrame,
) -> None:
    """
    Write the four result DataFrames to an Excel workbook in constant memory.

    Uses openpyxl write_only so the large CAZyme detail table is streamed row by
    row instead of being accumulated in an in-memory workbook (the previous
    memory risk). The workbook always contains the four sheets (an empty frame
    yields a one-row notice), and a per-sheet failure is isolated so it can never
    break the rest of the workbook. The four TSV outputs are unchanged.
    """
    try:
        from openpyxl import Workbook
    except Exception as exc:  # pragma: no cover - environment guard
        note = xlsx_path.with_suffix(xlsx_path.suffix + ".NOT_CREATED.txt")
        note.write_text(
            f"XLSX workbook was not created because openpyxl could not be imported: {exc}\n"
            f"All TSV outputs remain valid and complete.\n",
            encoding="utf-8",
        )
        print(f"[WARN] XLSX not created. See: {note}", flush=True)
        return

    sheets = [
        ("cazy_hits_detail",       detail_df,   "no CAZyme hits detected"),
        ("cazy_families_binary",   binary_df,   "no families passed filter"),
        ("cazy_summary_by_genome", summary_df,  "no summary data"),
        ("cazy_failures",          failures_df, "no failures"),
    ]

    wb = Workbook(write_only=True)
    for sheet_name, df, fallback in sheets:
        try:
            ws = wb.create_sheet(title=sheet_name[:31])
            _append_df_to_sheet(ws, df, fallback)
        except Exception as exc:
            try:
                ws_err = wb.create_sheet(title=(sheet_name[:27] + "_ERR"))
                ws_err.append(["write_error"])
                ws_err.append([str(exc)])
                print(f"[WARN] Sheet '{sheet_name}' could not be written normally: {exc}", flush=True)
            except Exception:
                pass
    wb.save(str(xlsx_path))


# ------------------------------------------------------------
# Empty-column schema
# ------------------------------------------------------------

DETAIL_COLS = [
    "genome_id", "input_faa", "protein_id", "protein_accession", "protein_length",
    "cazy_family", "family_accession", "family_length",
    "full_evalue", "full_score", "full_bias",
    "domain_number", "domain_of", "c_evalue", "i_evalue",
    "domain_score", "domain_bias",
    "hmm_from", "hmm_to", "ali_from", "ali_to", "env_from", "env_to",
    "acc", "hmm_coverage", "passed_filter",
]

FAILURE_COLS = ["genome_id", "input_faa", "status", "message"]


# ------------------------------------------------------------
# Main
# ------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description="CAZyme screening with hmmscan against the dbCAN HMM database."
    )
    parser.add_argument("--manifest",      required=True,  help="TSV manifest with genome_id and normalized_faa columns")
    parser.add_argument("--raw-dir",       required=True,  help="Directory for hmmscan .domtblout raw files")
    parser.add_argument("--output-tsv-dir", required=True, help="Output directory for TSV results")
    parser.add_argument("--output-xlsx-dir", required=True, help="Output directory for XLSX results")

    # Normal-mode arguments (not needed in --resume mode)
    parser.add_argument("--dbcan-db",      default=None,   help="Path to compiled dbCAN HMM database (hmmpress output)")
    parser.add_argument("--threads",       type=int, default=2)
    parser.add_argument("--hmmscan-bin",   default="hmmscan")

    # Resume mode
    parser.add_argument(
        "--resume",
        action="store_true",
        help="Skip hmmscan and re-parse existing .domtblout files from --raw-dir",
    )
    args = parser.parse_args()

    if not args.resume and args.dbcan_db is None:
        parser.error("--dbcan-db is required unless --resume is set")

    raw_dir        = Path(args.raw_dir)
    output_tsv_dir = Path(args.output_tsv_dir)
    output_xlsx_dir = Path(args.output_xlsx_dir)

    for d in (raw_dir, output_tsv_dir, output_xlsx_dir):
        ensure_dir(d)

    manifest_df = pd.read_csv(args.manifest, sep="\t")
    manifest_df["genome_id"] = manifest_df["genome_id"].astype(str)

    detail_frames = []
    failure_rows  = []

    for _, row in manifest_df.iterrows():
        genome_id = str(row["genome_id"])
        faa       = safe_text(row.get("normalized_faa", ""))
        domtblout = raw_dir / f"{genome_id}.dbcan.domtblout"

        print(f"\n=== {'[RESUME] Parsing' if args.resume else 'Processing'} genome: {genome_id} ===", flush=True)
        print(f"FAA      : {faa}", flush=True)
        print(f"domtblout: {domtblout}", flush=True)

        # --- Resume mode ---
        if args.resume:
            if not domtblout.exists():
                msg = f"Missing domtblout file: {domtblout}"
                failure_rows.append({"genome_id": genome_id, "input_faa": faa,
                                     "status": "FAILED", "message": msg})
                print(f"[ERROR] {msg}", flush=True)
                continue

        # --- Normal mode ---
        else:
            if not faa or not Path(faa).is_file():
                msg = "Missing normalized_faa file"
                failure_rows.append({"genome_id": genome_id, "input_faa": faa,
                                     "status": "FAILED", "message": msg})
                print(f"[SKIP] {msg} for genome: {genome_id}", flush=True)
                continue

            # Skip if .domtblout already exists (incremental mode)
            if domtblout.exists():
                print(f"[SKIP] domtblout already exists, skipping hmmscan for: {genome_id}", flush=True)
            else:
                cmd = [
                    args.hmmscan_bin, "--cpu", str(args.threads),
                    "--domtblout", str(domtblout),
                    str(args.dbcan_db), str(faa),
                ]
                print(f"[RUN] {' '.join(cmd)}", flush=True)
                result = run_command(cmd)

                if result.returncode != 0:
                    msg = f"hmmscan failed. STDERR:\n{result.stderr.strip()}"
                    failure_rows.append({"genome_id": genome_id, "input_faa": faa,
                                         "status": "FAILED", "message": msg})
                    print(f"[ERROR] hmmscan failed for genome: {genome_id}", flush=True)
                    if result.stderr.strip():
                        print(result.stderr.strip(), flush=True)
                    continue

        # --- Parse .domtblout (common to both modes) ---
        parsed = parse_domtblout(domtblout, genome_id, faa)
        if not parsed.empty:
            detail_frames.append(parsed)
        else:
            print(f"[INFO] No hits parsed for genome: {genome_id}", flush=True)

    # --- Aggregate results ---
    if detail_frames:
        detail_df = pd.concat(detail_frames, ignore_index=True)
    else:
        detail_df = pd.DataFrame(columns=DETAIL_COLS)

    binary_df   = build_binary_matrix(detail_df, manifest_df)
    summary_df  = build_summary(detail_df, manifest_df)
    failures_df = (
        pd.DataFrame(failure_rows, columns=FAILURE_COLS)
        if failure_rows
        else pd.DataFrame(columns=FAILURE_COLS)
    )

    # --- Write outputs ---
    detail_path  = output_tsv_dir / "cazy_hits_detail.tsv"
    binary_path  = output_tsv_dir / "cazy_families_binary.tsv"
    summary_path = output_tsv_dir / "cazy_summary_by_genome.tsv"
    failures_path = output_tsv_dir / "cazy_failures.tsv"
    xlsx_path    = output_xlsx_dir / "cazy_dbcan_results.xlsx"

    detail_df.to_csv(detail_path,   sep="\t", index=False, lineterminator="\n")
    binary_df.to_csv(binary_path,   sep="\t", index=False, lineterminator="\n")
    summary_df.to_csv(summary_path, sep="\t", index=False, lineterminator="\n")
    failures_df.to_csv(failures_path, sep="\t", index=False, lineterminator="\n")

    write_xlsx(xlsx_path, detail_df, binary_df, summary_df, failures_df)

    mode_label = "resume from raw" if args.resume else "full scan"
    print(f"\n[DONE] CAZyme screening completed ({mode_label}).", flush=True)
    print(f"Detail TSV    : {detail_path}",  flush=True)
    print(f"Binary TSV    : {binary_path}",  flush=True)
    print(f"Summary TSV   : {summary_path}", flush=True)
    print(f"Failures TSV  : {failures_path}", flush=True)
    print(f"Workbook XLSX : {xlsx_path}",    flush=True)


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"[FATAL] {exc}", file=sys.stderr)
        sys.exit(1)

